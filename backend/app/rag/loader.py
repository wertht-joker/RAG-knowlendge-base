"""文档解析：从 PDF / Word / Excel / Markdown 中提取文本"""
import pymupdf
from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook


def load_pdf(path: str) -> str:
    """解析 PDF，逐页提取文本"""
    texts = []
    with pymupdf.open(path) as doc:
        for page in doc:
            texts.append(page.get_text())
    return "\n".join(texts)


def load_docx(path: str) -> str:
    """解析 Word，按文档顺序提取段落与表格（表格转成「单元格 | 单元格」格式）"""
    doc = DocxDocument(path)
    parts = []
    # 遍历 body 子元素，保持段落与表格的先后顺序
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            parts.append(Paragraph(child, doc).text)
        elif child.tag == qn("w:tbl"):
            for row in Table(child, doc).rows:
                cells = [cell.text.strip() for cell in row.cells]
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def load_xlsx(path: str) -> str:
    """解析 Excel，遍历每个 Sheet，每行转成「单元格 | 单元格」格式"""
    wb = load_workbook(path, read_only=True, data_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"【{sheet.title}】")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            line = " | ".join(cells).rstrip(" |")
            if line.strip():
                parts.append(line)
    return "\n".join(parts)


def load_markdown(path: str) -> str:
    """解析 Markdown / 纯文本，直接读取"""
    with open(path, encoding="utf-8") as f:
        return f.read()


def load_document(path: str, file_type: str) -> str:
    """根据文件类型分发到对应解析器，返回纯文本"""
    if file_type == "pdf":
        return load_pdf(path)
    if file_type == "docx":
        return load_docx(path)
    if file_type == "xlsx":
        return load_xlsx(path)
    if file_type in ("md", "txt"):
        return load_markdown(path)
    raise ValueError(f"不支持的文件格式: {file_type}")
